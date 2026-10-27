"""Builds the CSCI 343 Milestone 1 document for Biohacker.

Produces:
  mockup_*.png                         labeled app mock-up boards
  <lastname>_CSCI343_milestone1.docx   one-page paper + mock-up (Calibri 12pt, 1.5 spacing)
  <lastname>_CSCI343_milestone1.pdf    same, exported with LibreOffice

Usage:  python3 make_milestone1.py "First Last" [YYYY-MM-DD]
"""
import os
import subprocess
import sys
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SCREENS = os.path.join(HERE, "screens")

NAVY = (13, 71, 161)
BLUE = (21, 101, 192)
SKY = (227, 242, 253)
TEXT = (11, 31, 58)
MUTED = (91, 107, 130)


def font(size, bold=False):
    candidates = [
        "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf" if bold else "/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
        "/Library/Fonts/Calibri.ttf",
    ]
    for c in candidates:
        if os.path.exists(c):
            return ImageFont.truetype(c, size)
    return ImageFont.load_default()


def phone(path, width):
    """Screenshot inside a simple phone frame."""
    shot = Image.open(path).convert("RGB")
    h = int(shot.height * width / shot.width)
    shot = shot.resize((width, h), Image.LANCZOS)
    bezel = 14
    frame = Image.new("RGB", (width + bezel * 2, h + bezel * 2), (255, 255, 255))
    d = ImageDraw.Draw(frame)
    d.rounded_rectangle([0, 0, frame.width - 1, frame.height - 1], radius=38, fill=TEXT)
    frame.paste(shot, (bezel, bezel))
    return frame


def wrap(draw, text, fnt, max_w):
    words, lines, line = text.split(), [], ""
    for w in words:
        test = f"{line} {w}".strip()
        if draw.textlength(test, font=fnt) <= max_w:
            line = test
        else:
            lines.append(line)
            line = w
    if line:
        lines.append(line)
    return lines


def board(title, subtitle, cells, out, cols=3):
    W = 2400
    col_w = (W - 120) // cols
    phone_w = col_w - 150
    rows = (len(cells) + cols - 1) // cols
    sample = phone(os.path.join(SCREENS, cells[0]["img"]), phone_w)
    cell_h = 90 + sample.height + 330
    H = 230 + rows * cell_h + 40
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W, 190], fill=NAVY)
    d.text((60, 40), title, font=font(72, True), fill=(255, 255, 255))
    d.text((60, 125), subtitle, font=font(38), fill=(187, 222, 251))

    for i, cell in enumerate(cells):
        r, c = divmod(i, cols)
        x0 = 60 + c * col_w
        y0 = 230 + r * cell_h
        # numbered label
        d.ellipse([x0, y0, x0 + 64, y0 + 64], fill=BLUE)
        num = str(cell["n"])
        d.text((x0 + 32 - d.textlength(num, font=font(40, True)) / 2, y0 + 8), num, font=font(40, True), fill=(255, 255, 255))
        d.text((x0 + 80, y0 + 2), cell["title"], font=font(44, True), fill=TEXT)
        d.text((x0 + 80, y0 + 50), cell["kind"], font=font(30), fill=BLUE)
        ph = phone(os.path.join(SCREENS, cell["img"]), phone_w)
        img.paste(ph, (x0 + (col_w - 60 - ph.width) // 2, y0 + 100))
        ty = y0 + 110 + ph.height + 10
        for comp in cell["components"]:
            for j, line in enumerate(wrap(d, comp, font(29), col_w - 100)):
                if j == 0:
                    d.ellipse([x0 + 8, ty + 12, x0 + 20, ty + 24], fill=BLUE)
                d.text((x0 + 32, ty), line, font=font(29), fill=TEXT)
                ty += 36
    img.save(out)
    return out


def nav_board(out):
    W, H = 2400, 1900
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W, 190], fill=NAVY)
    d.text((60, 40), "Biohacker – Navigation Map & Modals", font=font(72, True), fill=(255, 255, 255))
    d.text((60, 125), "NavigationContainer → Native Stack (root) → Bottom Tabs", font=font(38), fill=(187, 222, 251))

    def box(x, y, w, h, label, sub=None, fill=SKY, border=BLUE):
        d.rounded_rectangle([x, y, x + w, y + h], radius=24, fill=fill, outline=border, width=4)
        d.text((x + 24, y + 16), label, font=font(38, True), fill=TEXT)
        if sub:
            d.text((x + 24, y + 62), sub, font=font(28), fill=MUTED)

    def arrow(x1, y1, x2, y2):
        d.line([x1, y1, x2, y2], fill=BLUE, width=5)
        import math
        a = math.atan2(y2 - y1, x2 - x1)
        for s in (-0.45, 0.45):
            d.line([x2, y2, x2 - 26 * math.cos(a + s), y2 - 26 * math.sin(a + s)], fill=BLUE, width=5)

    box(60, 230, 1160, 110, "App.js", "SafeAreaProvider · Context Provider · Splash · Fonts · StatusBar · Biometric lock")
    box(60, 400, 1160, 110, "Root Native Stack Navigator", "Tabs + stack screens below")
    arrow(640, 340, 640, 400)
    tabs = [("1 Home", 60), ("2 Library", 300), ("7 Quizzes", 540), ("10 My Gym", 780), ("4 My Stack", 1020)]
    d.text((60, 560), "Bottom Tab Navigator", font=font(36, True), fill=BLUE)
    for label, x in tabs:
        box(x, 610, 210, 90, label)
    arrow(640, 510, 640, 560)
    stack = [("3 Compound Detail", 60, 800), ("5 Journal", 470, 800), ("6 Settings", 880, 800),
             ("8 Quiz Play", 60, 940), ("9 Quiz Result", 470, 940)]
    d.text((60, 750), "Stack screens (pushed on top of tabs)", font=font(36, True), fill=BLUE)
    for label, x, y in stack:
        box(x, y, 370, 90, label)
    d.text((60, 1080), "Modals (do not count as pages)", font=font(36, True), fill=BLUE)
    modals = [("M1 Dose Modal", "Add to stack / log dose · TextInput validation · Alert", 1130),
              ("M2 Equipment Shop", "FlatList of machines · level locks · cost", 1260),
              ("M3 Gym Design", "Gym name TextInput · floor & wall styles", 1390)]
    for label, sub, y in modals:
        box(60, y, 1160, 110, label, sub, fill=(255, 255, 255))

    flow = [
        "Learning → reward loop:",
        "Library (read compound info)",
        "→ Quizzes (questions built from that info)",
        "→ Result: grade A–F → cash + XP",
        "→ XP raises level → unlocks machines,",
        "   floor/wall styles, bigger floor",
        "→ Gym: buy & place machines →",
        "   customers → hourly revenue",
    ]
    d.rounded_rectangle([1290, 230, 2340, 760], radius=24, fill=SKY, outline=BLUE, width=4)
    y = 260
    for i, line in enumerate(flow):
        d.text((1320, y), line, font=font(38, i == 0), fill=TEXT)
        y += 60

    ph = phone(os.path.join(SCREENS, "dosemodal.png"), 330)
    img.paste(ph, (1300, 820))
    d.text((1300 + ph.width + 30, 840), "M1 Dose Modal", font=font(40, True), fill=TEXT)
    for i, t in enumerate(["Numeric TextInput", "Inline error message", "Slot chips (map)", "Alert if dose is", "above typical range"]):
        d.text((1300 + ph.width + 30, 900 + i * 44), "• " + t, font=font(30), fill=TEXT)
    land = Image.open(os.path.join(SCREENS, "gym-landscape.png")).convert("RGB")
    land = land.resize((1160, int(land.height * 1160 / land.width)), Image.LANCZOS)
    d.text((60, 1550), "Landscape orientation: the Gym re-flows into two columns", font=font(36, True), fill=BLUE)
    canvas = Image.new("RGB", (W, 1610 + land.height + 40), (255, 255, 255))
    canvas.paste(img.crop((0, 0, W, 1610)), (0, 0))
    canvas.paste(land, (60, 1610))
    img = canvas
    img.save(out)
    return out


SCREENS_A = [
    dict(n=1, title="Home", kind="Bottom tab", img="home.png", components=[
        "Gradient header, greeting, cash pill",
        "Level/XP card (links to Quizzes)",
        "Quick-link icon buttons",
        "Today's stack checklist (map)",
        "Compound of the day (useMemo)",
        "Linking to source compound tree"]),
    dict(n=2, title="Compound Library", kind="Bottom tab", img="library.png", components=[
        "Search TextInput with validation",
        "Family filter chips (horizontal map)",
        "FlatList of compound cards (1–3 cols)",
        "Favorite star toggle (icons)"]),
    dict(n=3, title="Compound Detail", kind="Stack screen", img="detail.png", components=[
        "useLayoutEffect sets header + favorite",
        "Web-loaded structure image (PubChem)",
        "HTTP fetch: formula & weight (useEffect)",
        "Benefits, dose, cautions sections",
        "Add to stack / Log dose → Dose Modal",
        "Linking: Examine, PubMed, PubChem"]),
    dict(n=4, title="My Stack", kind="Bottom tab", img="stack.png", components=[
        "Grouped by Morning / Pre-workout / Evening",
        "Edit (Dose Modal) & remove (Alert)",
        "'Log all' per time slot",
        "Saved to internal storage"]),
    dict(n=5, title="Journal", kind="Stack screen", img="journal.png", components=[
        "Stats cards (doses, days, most taken)",
        "FlatList grouped by day",
        "Delete entry with Alert confirm"]),
    dict(n=6, title="Settings", kind="Stack screen", img="settings.png", components=[
        "Name TextInput with validation",
        "Switches: haptics, biometric lock, cautions, compact",
        "Biometric (Face ID / fingerprint)",
        "Reset data (Alert), About + links"]),
]

SCREENS_B = [
    dict(n=7, title="Quizzes", kind="Bottom tab", img="quizzes.png", components=[
        "Level, completed, avg grade, cash",
        "XP progress bar",
        "FlatList of 11 quizzes + best grade"]),
    dict(n=8, title="Quiz Play", kind="Stack screen", img="quizplay.png", components=[
        "Animated progress bar",
        "Multiple-choice answers (map)",
        "Right/wrong feedback + explanation",
        "Haptics; Alert before quitting"]),
    dict(n=9, title="Quiz Result", kind="Stack screen", img="result.png", components=[
        "Animated grade badge A–F",
        "Cash + XP earned, level",
        "Level-up unlock list",
        "Go to gym / retake / all quizzes"]),
    dict(n=10, title="My Gym (3D)", kind="Bottom tab", img="gym.png", components=[
        "Isometric 3D gym (react-native-svg)",
        "Rotate camera 90°, tap tiles to build",
        "Animated customers walking",
        "Stats: customers, $/hr, rating",
        "Collect revenue; Move/Rotate/Sell"]),
    dict(n=11, title="Equipment Shop", kind="Modal (M2)", img="shop.png", components=[
        "15 machines unlock by level",
        "Cost, customers drawn, bonus income",
        "Alert if locked or too expensive"]),
    dict(n=12, title="Gym Design", kind="Modal (M3)", img="design.png", components=[
        "Gym name TextInput + validation",
        "Floor & wall style swatches",
        "Locked styles show required level"]),
]

PAPER = [
    ("Purpose", "Biohacker is a mobile app for people who want to improve their focus, fitness, sleep and long-term "
     "health with supplements and learn about them safely. The supplement world is "
     "full of hype, and beginners rarely know what a compound does, how much to take, or whether it is legal. Biohacker organizes information the same way as the Max Pharma compound tree (family › group › "
     "compound). It includes only substances that are legal to buy over the counter in the United States. Steroids, "
     "SARMs, research peptides and prescription drugs are left out on purpose."),
    ("Goal", "The goal of the app is to turn research into safe habits and real knowledge. Users browse 32 legal "
     "compounds with doses, evidence grades and cautions, build a daily stack, and log doses in a journal. To make learning stick, Biohacker adds a game. Quizzes are generated directly from the library, "
     "and each grade (A–F) earns cash and XP. XP raises the user’s level, which unlocks new equipment, floor and wall "
     "styles, and a larger floor in a virtual 3D gym. Users design the gym and place machines, "
     "which attract customers who produce hourly revenue. The meaningful result is a user who knows safe dosing and keeps a consistent routine."),
    ("Tools and Components", "Biohacker is built with React Native and Expo. Screens are organized with React "
     "Navigation, using a bottom tab bar for the five main sections and stacked screens for details, quizzes, the "
     "journal and settings, for ten pages plus three pop-up windows. The app has a navigation container, safe area "
     "support, a status bar, custom Poppins fonts, a splash screen, and a theme file that holds the blue and white "
     "colors. Layouts adjust automatically to any phone or tablet and to both portrait and landscape. The app uses "
     "text inputs with validation, links to outside websites, images loaded from the web, React hooks (useMemo, "
     "useLayoutEffect and useEffect), alert pop-ups, icons, switches, gradient backgrounds, modals, scrolling lists, "
     "and app-wide state management with the Context API. It requests live chemistry data from the free PubChem "
     "database. Device features include saving data to the phone’s storage, vibration feedback, and a Face ID or "
     "fingerprint lock. A community drawing library (react-native-svg) is used to draw the 3D gym."),
]


def set_font(run, size=12, bold=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold


def build_doc(name, when, boards, out_docx):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    style.font.size = Pt(12)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    for text, bold in (("Biohacker", True), (name, False), (when, False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), 12, bold)

    for heading, body in PAPER:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5)
        set_font(p.add_run(f"{heading}. "), 12, True)
        set_font(p.add_run(body), 12)

    for b in boards:
        doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img = Image.open(b)
        max_w, max_h = 6.5, 8.6
        w = min(max_w, max_h * img.width / img.height)
        p.add_run().add_picture(b, width=Inches(w))
    doc.save(out_docx)


def main():
    name = sys.argv[1] if len(sys.argv) > 1 else "Your Name"
    when = sys.argv[2] if len(sys.argv) > 2 else date.today().strftime("%B %-d, %Y")
    last = name.split()[-1].lower()
    boards = [
        board("Biohacker – App Mock-up (1/3)", "Core screens: learn, track and log", SCREENS_A, os.path.join(HERE, "mockup_1_core.png")),
        board("Biohacker – App Mock-up (2/3)", "Gamification: quizzes unlock the virtual 3D gym", SCREENS_B, os.path.join(HERE, "mockup_2_game.png")),
        nav_board(os.path.join(HERE, "mockup_3_navigation.png")),
    ]
    out_docx = os.path.join(HERE, f"{last}_CSCI343_milestone1.docx")
    build_doc(name, when, boards, out_docx)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", HERE, out_docx], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print("Wrote", out_docx, "and PDF")


if __name__ == "__main__":
    main()
